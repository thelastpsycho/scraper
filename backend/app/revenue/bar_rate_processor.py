"""Extract the BAR rate ladder from the fixed "Rate Structure" workbook.

Reads the "7. Rate Strategy" sheet's Publish / BAR1-BAR10 / MyValue rows (one
row per level, one column per room type) and stores them as a BAR-only SQLite
DB (``bar_rates.db``), independent of the scrape -> yield -> allotment
pipeline. The BAR Calculator frontend page reads this table to let a user
preview how editing a BAR level's discount % moves its rate off the Publish
rate for each room type.

The source workbook (``data/bar_rate_structure.xlsx``) is a fixed reference
document, not something re-uploaded per run - replace that file and re-run
this script when the rate structure itself changes (e.g. next year's budget).

The sheet layout (row/column positions) is fixed by convention in this
workbook rather than discovered generically - the label cells are checked
against their expected text so a changed template fails loudly instead of
silently importing the wrong numbers.
"""

import os
import re
import sqlite3

import openpyxl
from ..infrastructure.paths import get_data_dir

SHEET_NAME = '7. Rate Strategy'
LABEL_COL = 5             # column E: 'R. Type' / 'Publish' labels
HEADER_ROW = 10            # room-type names, columns F:R
PUBLISH_ROW = 11           # base "Publish" rate for each room type
BAR_LABEL_COL = 2          # column B, on the BAR-level rows only
DISCOUNT_COL = 5           # column E, on the BAR-level rows only
BAR_LEVEL_ROWS = range(16, 26)   # 'Internet Rate / WIG Rate / BAR 1' .. 'BAR 10'
MYVALUE_ROW = 26           # loyalty / MyValue rate (no discount %)
ROOM_TYPE_COL_START = 6    # column F
ROOM_TYPE_COL_END = 18     # column R (inclusive)


def process_bar_rates():
    data_dir = get_data_dir()
    source_path = os.path.join(data_dir, 'bar_rate_structure.xlsx')

    if not os.path.exists(source_path):
        raise FileNotFoundError(f"The fixed BAR rate structure Excel file was not found at: {source_path}")

    wb = openpyxl.load_workbook(source_path, data_only=True)
    if SHEET_NAME not in wb.sheetnames:
        raise ValueError(f"Sheet '{SHEET_NAME}' not found. Available sheets: {wb.sheetnames}")
    ws = wb[SHEET_NAME]

    header_label = ws.cell(row=HEADER_ROW, column=LABEL_COL).value
    if not header_label or 'r. type' not in str(header_label).strip().lower():
        raise ValueError(
            f"Expected 'R. Type' at {SHEET_NAME}!{ws.cell(row=HEADER_ROW, column=LABEL_COL).coordinate}, "
            f"found {header_label!r}. The Rate Strategy sheet layout may have changed."
        )

    room_types = []
    for col in range(ROOM_TYPE_COL_START, ROOM_TYPE_COL_END + 1):
        name = ws.cell(row=HEADER_ROW, column=col).value
        if name:
            room_types.append((col, str(name).strip()))
    if not room_types:
        raise ValueError(f"No room types found on row {HEADER_ROW} of '{SHEET_NAME}'")

    publish_label = ws.cell(row=PUBLISH_ROW, column=LABEL_COL).value
    if not publish_label or 'publish' not in str(publish_label).strip().lower():
        raise ValueError(
            f"Expected 'Publish' at {SHEET_NAME}!{ws.cell(row=PUBLISH_ROW, column=LABEL_COL).coordinate}, "
            f"found {publish_label!r}."
        )

    rows = []

    def add_level(row_idx, level_name, discount_pct):
        for col, room_type in room_types:
            rate = ws.cell(row=row_idx, column=col).value
            if rate is None:
                continue
            rows.append({
                'room_type': room_type,
                'level': level_name,
                'discount_pct': discount_pct,
                'rate': float(rate),
            })

    add_level(PUBLISH_ROW, 'Publish', None)

    for row_idx in BAR_LEVEL_ROWS:
        label = ws.cell(row=row_idx, column=BAR_LABEL_COL).value
        if not label:
            continue
        match = re.search(r'BAR\s*(\d+)', str(label), re.IGNORECASE)
        if not match:
            continue
        discount = ws.cell(row=row_idx, column=DISCOUNT_COL).value
        discount_pct = float(discount) * 100 if isinstance(discount, (int, float)) else None
        add_level(row_idx, f"BAR{match.group(1)}", discount_pct)

    if ws.cell(row=MYVALUE_ROW, column=BAR_LABEL_COL).value:
        add_level(MYVALUE_ROW, 'MyValue', None)

    if not rows:
        raise ValueError("No BAR rate rows were extracted - check the Rate Strategy sheet layout")

    db_path = os.path.join(data_dir, 'bar_rates.db')
    conn = sqlite3.connect(db_path)
    conn.execute('DROP TABLE IF EXISTS bar_rates')
    conn.execute('''
        CREATE TABLE bar_rates (
            room_type TEXT NOT NULL,
            level TEXT NOT NULL,
            discount_pct REAL,
            rate REAL NOT NULL,
            PRIMARY KEY (room_type, level)
        )
    ''')
    conn.executemany(
        'INSERT INTO bar_rates (room_type, level, discount_pct, rate) '
        'VALUES (:room_type, :level, :discount_pct, :rate)',
        rows
    )
    conn.commit()
    conn.close()

    message = f"Processed {len(rows)} BAR rate rows for {len(room_types)} room types"
    print(f'{message} into {db_path}')
    return message


if __name__ == '__main__':
    process_bar_rates()
